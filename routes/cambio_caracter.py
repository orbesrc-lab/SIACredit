import os
import io
import json
import uuid
import datetime
import PyPDF2
import docx2txt
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from flask import Blueprint, jsonify, request, render_template, send_file
from routes.ai import call_ai
from utils.db import supabase, get_active_inst_id
from utils.auth import require_permission
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

cambio_caracter_bp = Blueprint('cambio_caracter', __name__)

# Requisitos combinados de Ley 749 y Decreto 2038
REQ_PATH = os.path.join(os.path.dirname(__file__), 'requisitos_cc.json')
with open(REQ_PATH, 'r', encoding='utf-8') as f:
    REQUISITOS = json.load(f)

DEFAULTS_PATH = os.path.join(os.path.dirname(__file__), 'cc_defaults.json')

def load_defaults():
    if os.path.exists(DEFAULTS_PATH):
        with open(DEFAULTS_PATH, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}

def get_cc_project(inst_id):
    table_id = f"CC_PROJ_{inst_id}"
    project = None
    try:
        res = supabase.table('statistics').select('data_json').eq('table_id', table_id).execute()
        if res.data:
            project = json.loads(res.data[0]['data_json'])
    except Exception as e:
        print(f"[CC] Error fetch: {e}")
    
    defaults = load_defaults()

    if not project:
        project = {
            "inst_id": inst_id,
            "created_at": datetime.datetime.now().isoformat(),
            "updated_at": datetime.datetime.now().isoformat(),
            "fases": defaults.get("fases", []),
            "quincenas": defaults.get("quincenas", []),
            "frentes": defaults.get("frentes", []),
            "productos": defaults.get("productos", []),
            "riesgos": defaults.get("riesgos", []),
            "brechas": defaults.get("brechas", []),
            "requisitos": {
                req["id"]: {
                    "status": "Pendiente",
                    "notes": "",
                    "evidences": [],
                    "preguntas_6": {
                        "existe": False,
                        "vigente": False,
                        "aprobado": False,
                        "pertinente": False,
                        "aplica": False,
                        "evidencia": False
                    }
                } for req in REQUISITOS
            },
            "ai_history": []
        }
    else:
        # Merge missing sections
        for key in ["fases", "quincenas", "frentes", "productos", "riesgos", "brechas"]:
            if key not in project or not project[key]:
                project[key] = defaults.get(key, [])
                
        if "requisitos" not in project or not project["requisitos"]:
            project["requisitos"] = {
                req["id"]: {
                    "status": "Pendiente",
                    "notes": "",
                    "evidences": [],
                    "preguntas_6": {
                        "existe": False,
                        "vigente": False,
                        "aprobado": False,
                        "pertinente": False,
                        "aplica": False,
                        "evidencia": False
                    }
                } for req in REQUISITOS
            }
        else:
            # Ensure all requirements have the 6 questions structure
            for req in REQUISITOS:
                rid = req["id"]
                if rid not in project["requisitos"]:
                    project["requisitos"][rid] = {
                        "status": "Pendiente",
                        "notes": "",
                        "evidences": [],
                        "preguntas_6": {
                            "existe": False,
                            "vigente": False,
                            "aprobado": False,
                            "pertinente": False,
                            "aplica": False,
                            "evidencia": False
                        }
                    }
                elif "preguntas_6" not in project["requisitos"][rid]:
                    project["requisitos"][rid]["preguntas_6"] = {
                        "existe": False,
                        "vigente": False,
                        "aprobado": False,
                        "pertinente": False,
                        "aplica": False,
                        "evidencia": False
                    }

        if "ai_history" not in project:
            project["ai_history"] = []
            if project.get("ai_gap_analysis"):
                project["ai_history"].append({
                    "id": str(uuid.uuid4()),
                    "titulo": "Diagnóstico y Análisis de Brechas Base (MEN Ley 749)",
                    "fecha": project.get("created_at", datetime.datetime.now().isoformat()),
                    "tipo": "full_diagnosis",
                    "frente": "Institucional General",
                    "producto": "P2 Diagnóstico Institucional",
                    "autor": "Asesor Experto MEN (IA)",
                    "tokens_est": len(project.get("ai_gap_analysis", "")) // 4,
                    "contenido": project.get("ai_gap_analysis", "")
                })

    return project

def save_cc_project(inst_id, data):
    table_id = f"CC_PROJ_{inst_id}"
    data["updated_at"] = datetime.datetime.now().isoformat()
    data_json = json.dumps(data, ensure_ascii=False)
    try:
        check = supabase.table('statistics').select('id').eq('table_id', table_id).execute()
        if check.data:
            supabase.table('statistics').update({'data_json': data_json}).eq('table_id', table_id).execute()
        else:
            supabase.table('statistics').insert({
                'table_id': table_id,
                'inst_id': inst_id,
                'data_json': data_json
            }).execute()
        return True
    except Exception as e:
        print(f"[CC] Error save: {e}")
        return False

@cambio_caracter_bp.route('/cambio_caracter', methods=['GET'])
def render_cambio_caracter():
    inst_id = get_active_inst_id()
    return render_template('cambio_caracter.html', inst_id=inst_id, requisitos=REQUISITOS)

@cambio_caracter_bp.route('/api/cambio_caracter/data', methods=['GET', 'POST'])
def api_cambio_caracter_data():
    inst_id = request.args.get('inst_id', get_active_inst_id(), type=int)
    
    if request.method == 'POST':
        data = request.json
        if save_cc_project(inst_id, data):
            return jsonify({"status": "success"})
        return jsonify({"status": "error", "message": "Error saving data"}), 500
        
    project = get_cc_project(inst_id)
    return jsonify({"status": "success", "data": project})

@cambio_caracter_bp.route('/api/cambio_caracter/upload', methods=['POST'])
def api_cambio_caracter_upload():
    try:
        inst_id = request.form.get('inst_id', get_active_inst_id())
        req_id = request.form.get('req_id', 'general')
        
        # Soportar múltiples archivos o archivo individual
        files = request.files.getlist('files')
        if not files or len(files) == 0:
            single = request.files.get('file')
            if single:
                files = [single]
        
        if not files:
            return jsonify({"status": "error", "message": "Falta archivo"}), 400
            
        uploaded_results = []
        for file in files:
            if not file or not file.filename:
                continue
            ext = os.path.splitext(file.filename)[1].lower()
            new_filename = f"{uuid.uuid4()}{ext}"
            path = f"inst_{inst_id}/cambio_caracter/{req_id}/{new_filename}"
            
            file_bytes = file.read()
            file_options = {"content-type": file.content_type} if hasattr(file, 'content_type') else {}
            supabase.storage.from_('evidencias').upload(path, file_bytes, file_options=file_options)
            
            public_url = supabase.storage.from_('evidencias').get_public_url(path)
            
            size_kb = len(file_bytes) / 1024
            size_str = f"{size_kb:.1f} KB" if size_kb < 1024 else f"{(size_kb/1024):.1f} MB"
            
            uploaded_results.append({
                "url": public_url,
                "name": file.filename,
                "path": path,
                "size": size_str,
                "ext": ext
            })
            
        if not uploaded_results:
            return jsonify({"status": "error", "message": "No se procesaron archivos válidos"}), 400
            
        # Si fue un solo archivo, retornar campos individuales para compatibilidad hacia atrás
        first_item = uploaded_results[0]
        return jsonify({
            "status": "success",
            "url": first_item["url"],
            "name": first_item["name"],
            "path": first_item["path"],
            "size": first_item["size"],
            "files": uploaded_results
        })
    except Exception as e:
        print(f"[CC] Upload error: {e}")
        return jsonify({"status": "error", "message": str(e)}), 500

@cambio_caracter_bp.route('/api/cambio_caracter/delete_file', methods=['POST'])
def api_cambio_caracter_delete_file():
    try:
        path = request.json.get('path')
        if path:
            supabase.storage.from_('evidencias').remove([path])
        return jsonify({"status": "success"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@cambio_caracter_bp.route('/api/cambio_caracter/margy_consult', methods=['POST'])
def api_cc_margy_consult():
    """Asesoría experta de Margy IA sobre qué exige realmente el MEN para un requisito específico."""
    try:
        data = request.json or {}
        req_id = data.get('req_id')
        user_message = (data.get('message') or '').strip()
        history = data.get('history', [])
        
        req_item = next((r for r in REQUISITOS if r['id'] == req_id), None)
        if not req_item:
            return jsonify({"status": "error", "message": "Requisito no encontrado"}), 404
            
        system_prompt = f"""Eres Margy IA, la Consejera Técnica Senior y Consultora Experta en Aseguramiento de Calidad del MEN y CONACES para la institución educativa.
Estás asesorando a la institución en su proceso estratégico de Cambio de Carácter Académico a Institución Universitaria bajo la Ley 749 de 2002 y el Decreto 2038 de 2023.

Estás orientando sobre el REQUISITO NORMATIVO ESPECÍFICO:
- Código: {req_item.get('id')}
- Eje: {req_item.get('eje')}
- Factor: {req_item.get('factor')}
- Título: {req_item.get('title')}
- Descripción del Requisito: {req_item.get('desc')}
- Exigencia Documental MEN: {req_item.get('evidencia_requerida')}
- Marco Legal: {req_item.get('ley')}

DIRECTRICES DE RESPUESTA:
1. Responde de manera profesional, cercana, empática y con profundo rigor técnico y normativo del MEN.
2. Si el usuario solicita orientación general o inicial, estructura tu respuesta en los siguientes 4 bloques en Markdown con formato claro y viñetas:
   - 🎯 **¿Qué busca realmente el MEN con este requisito?**: Explica el sentido profundo y por qué los Pares Evaluadores lo exigen (más allá del texto frío de la norma).
   - 📑 **Evidencias Clave y Coherencia Documental**: Detalla qué documentos específicos (actas, resoluciones, matrices, diagnósticos, firmas) deben existir para que el MEN lo considere plenamente cumplido.
   - 📈 **Evolución Recomendada por Fase**: Cómo debe madurar este requisito a lo largo del proceso institucional:
     * *Fase I (Alineación / Levantamiento):* Diagnóstico inicial y mapeo de brechas.
     * *Fase II (Diagnóstico Integral y Elaboración):* Documento técnico y propuesta fundamentada.
     * *Fase III (Aprobación Colegiada):* Acta y Acuerdo formal de Consejo Directivo o Superior.
     * *Fase IV (Radicación SACES / Consolidación):* Inclusión en el Documento Maestro final.
   - ⚠️ **Riesgos y Errores Comunes**: Qué inconsistencias suelen detectar los Pares Académicos del MEN que originan requerimientos o conceptos negativos.
3. Si el usuario hace una pregunta puntual (ej. sobre actas, cronogramas, formatos), respóndele de forma concisa y práctica orientada a la solución.
4. Mantén siempre el avatar y personalidad de Margy IA: pedagógica, rigurosa, motivadora y experta en el MEN."""

        messages = [{"role": "system", "content": system_prompt}]
        for h in history[-4:]:
            messages.append({"role": h.get("role", "user"), "content": h.get("content", "")})
            
        if user_message:
            messages.append({"role": "user", "content": user_message})
        else:
            messages.append({"role": "user", "content": f"Margy, por favor explícame en detalle qué es lo que realmente busca el MEN con este requisito '{req_item.get('title')}' y cómo debemos construir y madurar las evidencias por fase."})

        answer = call_ai(messages, max_tokens=2000, temperature=0.35)
        return jsonify({"status": "success", "answer": answer, "req_item": req_item})
    except Exception as e:
        print(f"[CC Margy] Error: {e}")
        return jsonify({"status": "error", "message": str(e)}), 500

@cambio_caracter_bp.route('/api/cambio_caracter/skel_users', methods=['GET'])
def api_cc_skel_users():
    inst_id = get_active_inst_id()
    try:
        res = supabase.table('users').select('id, email, name, role, inst_id').execute()
        valid_users = [
            u for u in (res.data or [])
            if u.get('role') in ['inst_admin', 'lider', 'operativo', 'admin']
            and (u.get('inst_id') in [inst_id, 7, None] or u.get('role') == 'admin')
        ]
        return jsonify({"status": "success", "users": valid_users})
    except Exception as e:
        print(f"[CC] Error fetch users: {e}")
        return jsonify({"status": "error", "message": str(e)}), 500

@cambio_caracter_bp.route('/api/cambio_caracter/ai_gap_analysis', methods=['POST'])
def api_cc_ai_gap_analysis():
    try:
        inst_id = request.form.get('inst_id', get_active_inst_id(), type=int)
        mode = request.form.get('mode', 'full_diagnosis') # full_diagnosis, frente_diagnosis, draft_product, evaluate_evidence, solve_brecha, custom_prompt
        frente_name = request.form.get('frente_name', '')
        producto_codigo = request.form.get('producto_codigo', '')
        brecha_id = request.form.get('brecha_id', '')
        custom_instructions = request.form.get('custom_instructions', '')
        direct_text = request.form.get('text', '')
        
        file = request.files.get('file')
        file_text = ""
        saved_doc_path = None
        
        if file:
            ext = os.path.splitext(file.filename)[1].lower()
            file_bytes = file.read()
            new_filename = f"base_doc_{uuid.uuid4()}{ext}"
            saved_doc_path = f"inst_{inst_id}/cambio_caracter/base_documents/{new_filename}"
            file_options = {"content-type": file.content_type} if hasattr(file, 'content_type') else {}
            try:
                supabase.storage.from_('evidencias').upload(saved_doc_path, file_bytes, file_options=file_options)
            except Exception as up_e:
                print(f"[CC AI] Upload warning: {up_e}")

            if ext == '.pdf':
                reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        file_text += extracted + "\n"
            elif ext in ['.docx', '.doc']:
                file_text = docx2txt.process(io.BytesIO(file_bytes))
            else:
                return jsonify({"status": "error", "message": "Formato no soportado para análisis IA. Sube PDF o Word."}), 400

        text_to_process = (file_text.strip() or direct_text.strip())
        if not text_to_process and mode not in ['draft_product', 'solve_brecha', 'custom_prompt']:
            return jsonify({"status": "error", "message": "Debes subir un archivo o ingresar texto para el análisis."}), 400

        # Truncate text to avoid token exhaustion if huge
        if len(text_to_process) > 35000:
            text_to_process = text_to_process[:35000]

        project = get_cc_project(inst_id)

        system_prompt = (
            "Eres el Asesor Experto Principal del Ministerio de Educación Nacional de Colombia (MEN) y de CONACES/CESU, "
            "especialista de máximo nivel en cambio de carácter académico institucional (Ley 749 de 2002, Ley 30 de 1992, "
            "Decreto 2216 de 2003, Decreto 1075 de 2015, Decreto 1330 de 2019, Decreto 2038 de 2023, Decreto 0529 de 2024 "
            "y Acuerdo CESU 01 de 2025). "
            "Tu estilo de redacción es formal, altamente técnico, académico, jurídico y propositivo. "
            "No te limitas a opinar: redactas articulados normativos, capítulos del PEI/PDI, justificaciones curriculares, "
            "políticas institucionales y matrices con rigor impecable, adaptados a la realidad de una institución que pasa a "
            "Institución Universitaria."
        )

        title_entry = "Análisis IA"
        task_frente = frente_name or "General"
        task_product = producto_codigo or "Transversal"

        # PROMPT BUILDING BY MODE TO MINIMIZE TOKENS
        if mode == 'frente_diagnosis':
            # Filter requirements for this front
            rubrica_frente = []
            for r in REQUISITOS:
                if frente_name.lower() in (r.get('eje', '') + ' ' + r.get('factor', '') + ' ' + r.get('desc', '')).lower():
                    rubrica_frente.append(f"- ID: {r['id']} | FACTOR: {r['factor']}\n  ASPECTO: {r['desc']}\n  EVIDENCIA EXIGIDA: {r.get('evidencia_requerida','')}")
            
            if not rubrica_frente:
                rubrica_frente = [f"- ID: {r['id']} | {r['desc']}" for r in REQUISITOS[:8]]
            rubrica_str = "\n".join(rubrica_frente[:10])

            title_entry = f"Diagnóstico Focalizado: Frente {frente_name}"
            user_prompt = f"""
Has sido convocado para liderar el diagnóstico y redacción del frente: **{frente_name}**.
Objetivo: Evaluar el estado actual, detectar brechas frente a los estándares de Institución Universitaria del MEN y redactar la propuesta de mejoramiento o articulado normativo correspondiente.

REQUISITOS APLICABLES AL FRENTE ({frente_name}):
{rubrica_str}

INSTRUCCIONES ESPECÍFICAS DEL USUARIO:
{custom_instructions or 'Identifica las brechas y redacta el articulado y plan de acción correspondiente.'}

TEXTO / EVIDENCIA BASE:
{text_to_process}

Debes responder con un bloque de código JSON con esta estructura exacta:
```json
{{
  "markdown_report": "Aquí redacta el INFORME TÉCNICO Y LA PROPUESTA REDACTADA (articulados, políticas, planes de acción) en formato Markdown.",
  "evaluations": {{
    "req_X": {{
      "status": "En Construcción",
      "notes": "Justificación técnica del estado y recomendación.",
      "preguntas_6": {{
        "existe": true,
        "vigente": true,
        "aprobado": false,
        "pertinente": true,
        "aplica": false,
        "evidencia": false
      }}
    }}
  }}
}}
```
"""

        elif mode == 'draft_product':
            prod_info = next((p for p in project.get("productos", []) if p["codigo"] == producto_codigo), None)
            prod_name = prod_info["nombre"] if prod_info else f"Producto {producto_codigo}"
            prod_min = prod_info["contenido_minimo"] if prod_info else ""
            title_entry = f"Redacción Experta: {producto_codigo} - {prod_name}"
            
            user_prompt = f"""
Se te encomienda REDACTAR EL ENTREGABLE INSTITUCIONAL: **{producto_codigo}: {prod_name}**.
Contenido mínimo exigido por la metodología institucional:
{prod_min}

Directrices del proceso de cambio de carácter (Ley 749 de 2002 / Decreto 2038 de 2023):
- Incluye Considerandos jurídicos y de pertinencia académica.
- Estructura el documento con Títulos, Capítulos, Artículos y Parágrafos donde aplique.
- Incorpora disposiciones transitorias para proteger las cohortes de estudiantes y docentes en transición.
- Define indicadores y responsables.

INSTRUCCIONES ESPECÍFICAS:
{custom_instructions or 'Redacta el documento formal completo, listo para revisión del Comité Técnico y aprobación.'}

INSUMOS / BASE DOCUMENTAL (si aplica):
{text_to_process or 'No se adjuntó texto previo; redacta la propuesta estándar de máxima calidad institucional.'}

Devuelve un bloque JSON:
```json
{{
  "markdown_report": "Aquí redacta el DOCUMENTO COMPLETO en formato Markdown formal y riguroso.",
  "evaluations": {{}}
}}
```
"""

        elif mode == 'solve_brecha':
            brecha_info = next((b for b in project.get("brechas", []) if b["id"] == brecha_id), None)
            desc_b = brecha_info["descripcion"] if brecha_info else "Brecha institucional"
            frente_b = brecha_info["frente"] if brecha_info else "General"
            accion_b = brecha_info.get("accion_propuesta", "Actualizar") if brecha_info else "Actualizar"
            title_entry = f"Solución Técnica a Brecha: {brecha_id}"

            user_prompt = f"""
Se requiere una SOLUCIÓN TÉCNICA Y REDACCIÓN NORMATIVA EXPERTA para cerrar la siguiente brecha identificada en el plan de cambio de carácter:
- Código: {brecha_id}
- Frente: {frente_b}
- Tratamiento institucional propuesto: {accion_b} (según metodología: Conservar, Fortalecer, Actualizar, Transformar, Armonizar, Crear, Derogar)
- Descripción de la Brecha: {desc_b}
- Evidencia faltante: {brecha_info.get('evidencia_faltante', '') if brecha_info else ''}

INSTRUCCIONES ESPECÍFICAS:
{custom_instructions or 'Redacta la solución definitiva: propuesta de articulado o reforma, procedimiento operativo y evidencias a consolidar.'}

INSUMOS / CONTEXTO ADICIONAL:
{text_to_process}

Devuelve un bloque JSON:
```json
{{
  "markdown_report": "Redacta el informe de cierre de brecha, propuesta de texto legal/académico y protocolo de evidencias.",
  "evaluations": {{}}
}}
```
"""

        elif mode == 'evaluate_evidence':
            title_entry = f"Auditoría de Evidencias (6 Preguntas) - {frente_name or 'General'}"
            user_prompt = f"""
Realiza una AUDITORÍA METODOLÓGICA DE EVIDENCIAS aplicando rigurosamente la 'Metodología de las 6 Preguntas Institucionales' de INTENALCO / MEN:
1. ¿Existe? (Verificación de existencia del instrumento o proceso)
2. ¿Está vigente? (Vigencia temporal, sin derogatorias tácitas)
3. ¿Está formalmente aprobado? (Por el órgano competente: Consejo Superior / Directivo / Académico con acta o acuerdo)
4. ¿Es pertinente frente al nuevo carácter? (Alineado a Institución Universitaria según Ley 749 y Decreto 2038)
5. ¿Se aplica efectivamente? (Operatividad real en la vida institucional)
6. ¿Existe evidencia de aplicación y resultados? (Registros, actas, mediciones, trazabilidad)

DOCUMENTO / EVIDENCIAS REVISADAS:
{text_to_process}

INSTRUCCIONES ESPECÍFICAS:
{custom_instructions or 'Emite dictamen de auditoría para cada requisito evidenciado y califica cada una de las 6 preguntas.'}

Devuelve un bloque JSON:
```json
{{
  "markdown_report": "Dictamen de auditoría con hallazgos, recomendaciones y plan de subsanación.",
  "evaluations": {{
    "req_X": {{
      "status": "Completado",
      "notes": "Dictamen fundado.",
      "preguntas_6": {{
        "existe": true,
        "vigente": true,
        "aprobado": true,
        "pertinente": true,
        "aplica": true,
        "evidencia": true
      }}
    }}
  }}
}}
```
"""

        else: # full_diagnosis (default)
            rubrica_text = "\n".join([f"- ID: {r['id']} | EJE: {r['eje']} | FACTOR: {r['factor']}\n  ASPECTO: {r['desc']}" for r in REQUISITOS[:25]])
            title_entry = "Diagnóstico Integral de Brechas y Propuesta de Redacción"
            user_prompt = f"""
La institución adelanta su proceso de cambio de carácter a 'Institución Universitaria' según la Ley 749 de 2002 y el Decreto 2038 de 2023.
A continuación te presento el documento base institucional.

Tu tarea es:
1. Identificar las brechas críticas de transformación frente a los estándares de Institución Universitaria.
2. Aplicar la matriz de tratamiento (Conservar, Fortalecer, Actualizar, Transformar, Armonizar, Crear, Derogar).
3. REDACTAR las propuestas de articulado normativo, ajuste de políticas y capítulos requeridos para subsanar los vacíos.

RUBRICA DE REFERENCIA (Factores Priorizados):
{rubrica_text}

INSTRUCCIONES ESPECÍFICAS:
{custom_instructions or 'Genera un informe integral con propuestas de redacción listas para usar en el Dossier Técnico.'}

DOCUMENTO BASE EXTRAÍDO:
{text_to_process}

Devuelve EXCLUSIVAMENTE un bloque de código JSON:
```json
{{
  "markdown_report": "Redacta el informe técnico de brechas y las propuestas de redacción formal de estatutos/políticas en formato Markdown.",
  "evaluations": {{
    "req_1": {{
      "status": "En Construcción",
      "notes": "Breve nota de cumplimiento.",
      "preguntas_6": {{
        "existe": true,
        "vigente": true,
        "aprobado": false,
        "pertinente": true,
        "aplica": false,
        "evidencia": false
      }}
    }}
  }}
}}
```
"""

        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ]

        ai_response = call_ai(messages, max_tokens=32000, inst_id=inst_id)

        import re
        match = re.search(r'```json\n(.*?)\n```', ai_response, re.DOTALL)
        if match:
            json_str = match.group(1)
        else:
            json_str = ai_response.strip()
            if json_str.startswith('```') and json_str.endswith('```'):
                json_str = json_str.strip('`').replace('json\n', '', 1).strip()

        try:
            parsed_data = json.loads(json_str)
            analysis_html = parsed_data.get("markdown_report", "")
            evaluations = parsed_data.get("evaluations", {})
        except Exception as json_e:
            print("[CC AI] JSON Parse fallback:", json_e)
            rescate = re.search(r'"markdown_report"\s*:\s*"(.*?)(?:","evaluations"|$)', json_str, re.DOTALL | re.IGNORECASE)
            if rescate:
                analysis_html = rescate.group(1).replace('\\n', '\n').replace('\\"', '"')
            else:
                analysis_html = ai_response
            evaluations = {}

        # Save to history entry
        tokens_est = (len(user_prompt) + len(analysis_html)) // 4
        history_id = str(uuid.uuid4())
        new_hist_item = {
            "id": history_id,
            "titulo": title_entry,
            "fecha": datetime.datetime.now().isoformat(),
            "tipo": mode,
            "frente": task_frente,
            "producto": task_product,
            "autor": "Asesor Experto MEN (IA)",
            "tokens_est": tokens_est,
            "contenido": analysis_html,
            "saved_doc_path": saved_doc_path
        }

        if "ai_history" not in project:
            project["ai_history"] = []
        project["ai_history"].insert(0, new_hist_item)
        project["ai_gap_analysis"] = analysis_html # Keep latest for backward compatibility

        # Update evaluations in project if any
        if evaluations:
            for req_id, eval_data in evaluations.items():
                if req_id in project["requisitos"]:
                    if eval_data.get("status") in ["Pendiente", "En Construcción", "Completado"]:
                        project["requisitos"][req_id]["status"] = eval_data["status"]
                    if eval_data.get("notes"):
                        project["requisitos"][req_id]["notes"] = eval_data["notes"]
                    if eval_data.get("preguntas_6"):
                        project["requisitos"][req_id]["preguntas_6"] = eval_data["preguntas_6"]

        # If it was a brecha solution, update the brecha item directly
        if mode == 'solve_brecha' and brecha_id:
            for b in project.get("brechas", []):
                if b["id"] == brecha_id:
                    b["solucion_redactada"] = analysis_html
                    b["estado"] = "Cerrada" if "conforme" in analysis_html.lower() else "En gestión"

        # If it was a product draft, update the product document content
        if mode == 'draft_product' and producto_codigo:
            for p in project.get("productos", []):
                if p["codigo"] == producto_codigo:
                    p["documento_contenido"] = analysis_html
                    p["estado"] = "En Construcción"

        save_cc_project(inst_id, project)

        return jsonify({
            "status": "success",
            "analysis_html": analysis_html,
            "evaluations": evaluations,
            "history_item": new_hist_item,
            "history_id": history_id,
            "tokens_est": tokens_est
        })

    except Exception as e:
        print(f"[CC AI] Error: {e}")
        return jsonify({"status": "error", "message": str(e)}), 500

@cambio_caracter_bp.route('/api/cambio_caracter/ai_history/save', methods=['POST'])
def api_cc_ai_history_save():
    try:
        inst_id = request.json.get('inst_id', get_active_inst_id())
        item_id = request.json.get('id')
        titulo = request.json.get('titulo', 'Documento Editado')
        contenido = request.json.get('contenido', '')
        frente = request.json.get('frente', 'General')
        producto = request.json.get('producto', 'Transversal')
        create_new = request.json.get('create_new', False)

        project = get_cc_project(inst_id)
        if "ai_history" not in project:
            project["ai_history"] = []

        if not create_new and item_id:
            # Update existing
            found = False
            for item in project["ai_history"]:
                if item["id"] == item_id:
                    item["titulo"] = titulo
                    item["contenido"] = contenido
                    item["frente"] = frente
                    item["producto"] = producto
                    item["fecha_modificacion"] = datetime.datetime.now().isoformat()
                    item["autor"] = "Usuario Institucional (Editado sin consumo de tokens)"
                    found = True
                    break
            if not found:
                item_id = str(uuid.uuid4())
                project["ai_history"].insert(0, {
                    "id": item_id,
                    "titulo": titulo,
                    "fecha": datetime.datetime.now().isoformat(),
                    "tipo": "manual_edit",
                    "frente": frente,
                    "producto": producto,
                    "autor": "Usuario Institucional",
                    "tokens_est": 0,
                    "contenido": contenido
                })
        else:
            # Create new version
            item_id = str(uuid.uuid4())
            project["ai_history"].insert(0, {
                "id": item_id,
                "titulo": titulo,
                "fecha": datetime.datetime.now().isoformat(),
                "tipo": "manual_edit",
                "frente": frente,
                "producto": producto,
                "autor": "Usuario Institucional (Versión guardada)",
                "tokens_est": 0,
                "contenido": contenido
            })

        save_cc_project(inst_id, project)
        return jsonify({"status": "success", "id": item_id, "message": "Guardado exitosamente en el histórico."})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@cambio_caracter_bp.route('/api/cambio_caracter/ai_history/delete', methods=['POST'])
def api_cc_ai_history_delete():
    try:
        inst_id = request.json.get('inst_id', get_active_inst_id())
        item_id = request.json.get('id')
        project = get_cc_project(inst_id)
        if "ai_history" in project:
            project["ai_history"] = [h for h in project["ai_history"] if h["id"] != item_id]
            save_cc_project(inst_id, project)
        return jsonify({"status": "success"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@cambio_caracter_bp.route('/api/cambio_caracter/download_word', methods=['POST'])
def api_cc_download_word():
    try:
        title = request.form.get('title', 'Documento Institucional - Cambio de Carácter')
        text = request.form.get('content', '')
        
        doc = docx.Document()
        
        # Page Margins (Normal 1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Cover / Header Style
        title_para = doc.add_paragraph()
        title_run = title_para.add_run("INTENALCO EDUCACIÓN SUPERIOR\nPROCESO DE CAMBIO DE CARÁCTER A INSTITUCIÓN UNIVERSITARIA")
        title_run.font.name = 'Calibri'
        title_run.font.size = Pt(11)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(99, 102, 241) # Indigo
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc_heading = doc.add_heading(title, level=0)
        doc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        date_para = doc.add_paragraph()
        d_run = date_para.add_run(f"Generado el: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} | Sistema SIAC")
        d_run.font.name = 'Calibri'
        d_run.font.size = Pt(9)
        d_run.font.italic = True
        d_run.font.color.rgb = RGBColor(148, 163, 184)
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # Parse markdown text line by line
        for line in text.split('\n'):
            line_str = line.strip()
            if not line_str:
                doc.add_paragraph()
                continue
            
            if line_str.startswith('# '):
                p = doc.add_heading(line_str[2:].replace('**', ''), level=1)
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
            elif line_str.startswith('## '):
                p = doc.add_heading(line_str[3:].replace('**', ''), level=2)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
            elif line_str.startswith('### '):
                p = doc.add_heading(line_str[4:].replace('**', ''), level=3)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
            elif line_str.startswith('- ') or line_str.startswith('* '):
                bullet_text = line_str[2:].replace('**', '')
                doc.add_paragraph(bullet_text, style='List Bullet')
            elif line_str.startswith('1. ') or line_str.startswith('2. ') or line_str.startswith('3. '):
                num_text = line_str[3:].replace('**', '')
                doc.add_paragraph(num_text, style='List Number')
            else:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6)
                # Bold markdown parsing (**text**)
                parts = line_str.split('**')
                for i, part in enumerate(parts):
                    if not part: continue
                    run = p.add_run(part)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(10.5)
                    if i % 2 == 1:
                        run.font.bold = True
                
        io_stream = io.BytesIO()
        doc.save(io_stream)
        io_stream.seek(0)
        safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '_', '-')).rstrip()
        filename = f"{safe_title[:40].replace(' ', '_')}_{datetime.datetime.now().strftime('%Y%m%d')}.docx"
        return send_file(io_stream, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as e:
        print(f"[CC] Word Download Error: {e}")
        return str(e), 500

@cambio_caracter_bp.route('/api/cambio_caracter/download_pdf', methods=['POST'])
def api_cc_download_pdf():
    try:
        title = request.form.get('title', 'Documento Institucional - Cambio de Caracter')
        text = request.form.get('content', '')
        from fpdf import FPDF
        
        class CC_PDF(FPDF):
            def header(self):
                self.set_font("helvetica", "B", 10)
                self.set_text_color(99, 102, 241)
                self.cell(0, 8, "INTENALCO - CAMBIO DE CARACTER A INSTITUCION UNIVERSITARIA", border=0, align="C")
                self.ln(5)
                self.set_font("helvetica", "I", 8)
                self.set_text_color(148, 163, 184)
                self.cell(0, 5, "Sistema de Aseguramiento de la Calidad (SIAC)", border=0, align="C")
                self.ln(8)
                self.set_draw_color(226, 232, 240)
                self.line(10, self.get_y(), 200, self.get_y())
                self.ln(6)
                
            def footer(self):
                self.set_y(-15)
                self.set_font("helvetica", "I", 8)
                self.set_text_color(148, 163, 184)
                self.cell(0, 10, f"Pagina {self.page_no()}", border=0, align="C")

        pdf = CC_PDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        
        # Document Title
        pdf.set_font("helvetica", "B", 14)
        pdf.set_text_color(30, 41, 59)
        safe_title = title.encode('latin-1', 'replace').decode('latin-1')
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 8, txt=safe_title, align="C")
        pdf.ln(5)
        
        pdf.set_font("helvetica", size=9)
        pdf.set_text_color(51, 65, 85)
        
        for raw_line in text.split('\n'):
            line = raw_line.strip()
            if not line:
                pdf.ln(4)
                continue
                
            pdf.set_x(pdf.l_margin)
            if line.startswith('# '):
                pdf.set_font("helvetica", "B", 12)
                pdf.set_text_color(67, 56, 202)
                safe_l = line[2:].replace('**', '').encode('latin-1', 'replace').decode('latin-1')
                pdf.ln(4)
                pdf.multi_cell(pdf.epw, 7, txt=safe_l)
                pdf.set_font("helvetica", size=9)
                pdf.set_text_color(51, 65, 85)
            elif line.startswith('## '):
                pdf.set_font("helvetica", "B", 10.5)
                pdf.set_text_color(79, 70, 229)
                safe_l = line[3:].replace('**', '').encode('latin-1', 'replace').decode('latin-1')
                pdf.ln(3)
                pdf.multi_cell(pdf.epw, 6, txt=safe_l)
                pdf.set_font("helvetica", size=9)
                pdf.set_text_color(51, 65, 85)
            elif line.startswith('### '):
                pdf.set_font("helvetica", "B", 9.5)
                pdf.set_text_color(30, 41, 59)
                safe_l = line[4:].replace('**', '').encode('latin-1', 'replace').decode('latin-1')
                pdf.ln(2)
                pdf.multi_cell(pdf.epw, 5, txt=safe_l)
                pdf.set_font("helvetica", size=9)
                pdf.set_text_color(51, 65, 85)
            else:
                safe_l = line.replace('**', '').encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(pdf.epw, 5, txt=safe_l)

        pdf_bytes = pdf.output()
        if isinstance(pdf_bytes, str):
            pdf_bytes = pdf_bytes.encode('latin-1')
            
        io_stream = io.BytesIO(pdf_bytes)
        filename = f"Documento_CC_{datetime.datetime.now().strftime('%Y%m%d')}.pdf"
        return send_file(io_stream, as_attachment=True, download_name=filename, mimetype="application/pdf")
    except Exception as e:
        print(f"[CC] PDF Download Error: {e}")
        return str(e), 500

@cambio_caracter_bp.route('/api/cambio_caracter/export_excel', methods=['GET'])
def api_cc_export_excel():
    try:
        inst_id = request.args.get('inst_id', get_active_inst_id(), type=int)
        project = get_cc_project(inst_id)

        wb = openpyxl.Workbook()
        
        # Styles
        header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid") # Dark Indigo
        sub_fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10)
        font_bold = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        def style_header_row(ws, cols):
            for col_num in range(1, cols + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = font_header
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 28

        def autofit_cols(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    if '\n' in val_str:
                        val_str = max(val_str.split('\n'), key=len)
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

        # 1. SHEET: FASES Y PLAN
        ws_fases = wb.active
        ws_fases.title = "Fases del Plan"
        headers_fases = ["Fase", "Nombre", "Periodo", "Propósito", "Productos e Hitos", "Progreso (%)", "Estado"]
        ws_fases.append(headers_fases)
        style_header_row(ws_fases, len(headers_fases))
        for row_idx, f in enumerate(project.get("fases", []), start=2):
            ws_fases.append([
                f.get("numero", ""),
                f.get("nombre", ""),
                f.get("periodo", ""),
                f.get("proposito", ""),
                f.get("productos_hitos", ""),
                f.get("progreso", 0),
                f.get("estado", "")
            ])
            for c in range(1, len(headers_fases) + 1):
                cell = ws_fases.cell(row=row_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_idx % 2 == 0: cell.fill = zebra_fill
        autofit_cols(ws_fases)

        # 2. SHEET: CRONOGRAMA QUINCENAL
        ws_q = wb.create_sheet(title="Cronograma Quincenal")
        headers_q = ["Quincena", "Hito Verificable", "Producto Asociado", "Responsable Líder", "Estado", "Notas"]
        ws_q.append(headers_q)
        style_header_row(ws_q, len(headers_q))
        for row_idx, q in enumerate(project.get("quincenas", []), start=2):
            ws_q.append([
                q.get("quincena", ""),
                q.get("hito", ""),
                q.get("producto_asociado", ""),
                q.get("responsable", ""),
                q.get("estado", ""),
                q.get("notas", "")
            ])
            for c in range(1, len(headers_q) + 1):
                cell = ws_q.cell(row=row_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_idx % 2 == 0: cell.fill = zebra_fill
        autofit_cols(ws_q)

        # 3. SHEET: FRENTES Y LÍDERES
        ws_fr = wb.create_sheet(title="Frentes y Líderes")
        headers_fr = ["Frente de Trabajo", "Líder Responsable", "Qué se Diagnosticará", "Qué se Hará (Hallazgos)", "Asesoría Experta", "Periodo", "Progreso (%)", "Estado"]
        ws_fr.append(headers_fr)
        style_header_row(ws_fr, len(headers_fr))
        for row_idx, fr in enumerate(project.get("frentes", []), start=2):
            ws_fr.append([
                fr.get("nombre", ""),
                fr.get("lider", ""),
                fr.get("que_diagnosticara", ""),
                fr.get("que_se_hara", ""),
                fr.get("asesoria", ""),
                fr.get("periodo", ""),
                fr.get("progreso", 0),
                fr.get("estado", "")
            ])
            for c in range(1, len(headers_fr) + 1):
                cell = ws_fr.cell(row=row_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_idx % 2 == 0: cell.fill = zebra_fill
        autofit_cols(ws_fr)

        # 4. SHEET: MATRIZ DE BRECHAS
        ws_br = wb.create_sheet(title="Matriz de Brechas")
        headers_br = ["Código", "Frente", "Fuente Normativa", "Descripción Brecha", "Evidencia Disponible", "Evidencia Faltante", "Tratamiento / Acción", "Responsable", "Plazo", "Prioridad", "Estado"]
        ws_br.append(headers_br)
        style_header_row(ws_br, len(headers_br))
        for row_idx, br in enumerate(project.get("brechas", []), start=2):
            ws_br.append([
                br.get("id", ""),
                br.get("frente", ""),
                br.get("fuente_normativa", ""),
                br.get("descripcion", ""),
                br.get("evidencia_disponible", ""),
                br.get("evidencia_faltante", ""),
                br.get("accion_propuesta", ""),
                br.get("responsable", ""),
                br.get("plazo", ""),
                br.get("prioridad", ""),
                br.get("estado", "")
            ])
            for c in range(1, len(headers_br) + 1):
                cell = ws_br.cell(row=row_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_idx % 2 == 0: cell.fill = zebra_fill
        autofit_cols(ws_br)

        # 5. SHEET: PRODUCTOS P1-P10
        ws_pr = wb.create_sheet(title="Productos Entregables")
        headers_pr = ["Código", "Producto", "Contenido Mínimo", "Fase Entrega", "Líder Responsable", "Versión", "Fecha Límite", "Estado"]
        ws_pr.append(headers_pr)
        style_header_row(ws_pr, len(headers_pr))
        for row_idx, pr in enumerate(project.get("productos", []), start=2):
            ws_pr.append([
                pr.get("codigo", ""),
                pr.get("nombre", ""),
                pr.get("contenido_minimo", ""),
                pr.get("fase", ""),
                pr.get("responsable", ""),
                pr.get("version", ""),
                pr.get("fecha_limite", ""),
                pr.get("estado", "")
            ])
            for c in range(1, len(headers_pr) + 1):
                cell = ws_pr.cell(row=row_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_idx % 2 == 0: cell.fill = zebra_fill
        autofit_cols(ws_pr)

        # 6. SHEET: RIESGOS Y CONTROLES
        ws_rk = wb.create_sheet(title="Matriz de Riesgos")
        headers_rk = ["Código", "Categoría", "Riesgo Identificado", "Probabilidad", "Impacto", "Control Preventivo / Mitigación", "Responsable", "Estado"]
        ws_rk.append(headers_rk)
        style_header_row(ws_rk, len(headers_rk))
        for row_idx, rk in enumerate(project.get("riesgos", []), start=2):
            ws_rk.append([
                rk.get("id", ""),
                rk.get("categoria", ""),
                rk.get("descripcion", ""),
                rk.get("probabilidad", ""),
                rk.get("impacto", ""),
                rk.get("control", ""),
                rk.get("responsable", ""),
                rk.get("estado", "")
            ])
            for c in range(1, len(headers_rk) + 1):
                cell = ws_rk.cell(row=row_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_idx % 2 == 0: cell.fill = zebra_fill
        autofit_cols(ws_rk)

        # 7. SHEET: REQUISITOS NORMATIVOS
        ws_rq = wb.create_sheet(title="Requisitos Normativos")
        headers_rq = ["ID", "Eje", "Factor", "Título", "Aspecto Exigido", "Evidencia Exigida", "Norma / Ley", "Estado", "Observaciones", "Evidencias Cargadas"]
        ws_rq.append(headers_rq)
        style_header_row(ws_rq, len(headers_rq))
        req_states = project.get("requisitos", {})
        for row_idx, req in enumerate(REQUISITOS, start=2):
            r_data = req_states.get(req["id"], {})
            ev_count = len(r_data.get("evidences", []))
            ws_rq.append([
                req.get("id", ""),
                req.get("eje", ""),
                req.get("factor", ""),
                req.get("title", ""),
                req.get("desc", ""),
                req.get("evidencia_requerida", ""),
                req.get("ley", ""),
                r_data.get("status", "Pendiente"),
                r_data.get("notes", ""),
                ev_count
            ])
            for c in range(1, len(headers_rq) + 1):
                cell = ws_rq.cell(row=row_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_idx % 2 == 0: cell.fill = zebra_fill
        autofit_cols(ws_rq)

        output_stream = io.BytesIO()
        wb.save(output_stream)
        output_stream.seek(0)
        filename = f"Matriz_Planeacion_Cambio_Caracter_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"
        return send_file(output_stream, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        print(f"[CC] Excel Export Error: {e}")
        return str(e), 500
